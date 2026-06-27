# -*- coding: utf-8 -*-
"""
report_generator.py — 產出 Word 報告（含補填表格）與 PDF 報告
"""

import io
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from config import TC4_THRESHOLD, OUTPUT_DIR


# ─── 樣式工具 ────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _bold_run(para, text: str, size_pt: int = 12, color: str = None):
    run      = para.add_run(text)
    run.bold = True
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


RESULT_COLORS = {
    "已完成":   "D4EDDA",
    "部分完成": "FFF3CD",
    "未完成":   "F8D7DA",
    "需補件":   "D1ECF1",
    "需確認":   "E2D9F3",
    "不適用":   "E9ECEF",
}

RESULT_EMOJI = {
    "已完成":   "✅",
    "部分完成": "⚠️",
    "未完成":   "❌",
    "需補件":   "📋",
    "需確認":   "❓",
    "不適用":   "—",
}


# ─── Word 報告 ────────────────────────────────────────────────────────

def generate_word_report(check_result: dict) -> bytes:
    doc = Document()

    # ── 封面 ──
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _bold_run(title_para, "專案規劃檢核報告", size_pt=18, color="1A237E")

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"檢核日期：{check_result.get('timestamp','')[:10]}　　"
        f"Member：{check_result.get('member_name','（未填）')}　　"
        f"檔案：{check_result.get('filename','')}"
    ).font.size = Pt(10)

    doc.add_paragraph()

    # ── 場景與評分摘要 ──
    doc.add_heading("📌 場景與評分摘要", level=1)

    score   = check_result.get("score", 0)
    passed  = check_result.get("passed", False)
    scene   = check_result.get("scene", "")
    level   = check_result.get("level", "")
    note    = check_result.get("scene_note", "")
    stats   = check_result.get("stats", {})
    comment = check_result.get("overall_comment", "")

    summary_tbl = doc.add_table(rows=5, cols=2)
    summary_tbl.style = "Table Grid"
    rows_data = [
        ("使用場景", scene),
        ("專案層級", f"{level} 級"),
        ("TC4 評分", f"{score} / 100"),
        ("開會門檻", "✅ 達標，可找 William 開會" if passed else f"❌ 未達標（需 {TC4_THRESHOLD} 分）"),
        ("場景備注", note or "—"),
    ]
    for i, (k, v) in enumerate(rows_data):
        summary_tbl.rows[i].cells[0].text = k
        summary_tbl.rows[i].cells[1].text = v
        _set_cell_bg(summary_tbl.rows[i].cells[0], "EEF2FF")

    doc.add_paragraph()

    # 統計列
    stat_para = doc.add_paragraph()
    for k, emoji in RESULT_EMOJI.items():
        stat_para.add_run(f"{emoji} {k}：{stats.get(k,0)} 項　")

    if comment:
        doc.add_paragraph()
        doc.add_heading("💬 整體評語", level=2)
        doc.add_paragraph(comment)

    # ── 逐項結果 ──
    doc.add_heading("📋 逐項檢核結果", level=1)

    items = check_result.get("items", [])
    need_fill = []

    for it in items:
        result  = it.get("result", "")
        emoji   = RESULT_EMOJI.get(result, "")
        color   = RESULT_COLORS.get(result, "FFFFFF")

        heading = doc.add_paragraph()
        _bold_run(heading, f"{emoji} 【序{it.get('seq','')}】{it.get('item','')}", size_pt=11)

        tbl = doc.add_table(rows=3, cols=2)
        tbl.style = "Table Grid"
        data = [
            ("判斷結果", result),
            ("文件依據", it.get("evidence", "文件未提及")),
            ("常見錯誤偵測", it.get("common_error_found", "—") or "—"),
        ]
        for i, (k, v) in enumerate(data):
            tbl.rows[i].cells[0].text = k
            tbl.rows[i].cells[1].text = v
            _set_cell_bg(tbl.rows[i].cells[0], "F5F5F5")
        _set_cell_bg(tbl.rows[0].cells[1], color)

        if it.get("suggestion"):
            sug_para = doc.add_paragraph()
            sug_para.add_run("💡 建議：").bold = True
            sug_para.add_run(it["suggestion"])

        if result in ("需補件", "部分完成"):
            need_fill.append(it)

        doc.add_paragraph()

    # ── 補填表格區 ──
    if need_fill:
        doc.add_page_break()
        doc.add_heading("📝 補件清單（請 Member 填寫後重新上傳）", level=1)
        doc.add_paragraph(
            f"以下 {len(need_fill)} 項需要補充資料。"
            "填寫完成後，請重新上傳此文件進行複檢。\n"
            f"⚠️ 目前評分 {score} 分，"
            + ("距開會門檻尚差 " + str(round(TC4_THRESHOLD - score, 1)) + " 分。" if not passed else "已達開會門檻。")
        )

        for it in need_fill:
            doc.add_paragraph()
            fill_title = doc.add_paragraph()
            _bold_run(fill_title, f"【補填 — 序{it.get('seq','')}：{it.get('item','')}】", color="1A237E")

            missing = it.get("missing", [])
            if not missing:
                missing = ["（請依「建議」欄說明補充）"]

            fill_tbl = doc.add_table(rows=len(missing) + 1, cols=3)
            fill_tbl.style = "Table Grid"

            # 標題列
            headers = fill_tbl.rows[0].cells
            headers[0].text = "需補充的資訊"
            headers[1].text = "說明 / AI 範例"
            headers[2].text = "請填入"
            for h in headers:
                _set_cell_bg(h, "1A237E")
                for run in h.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.bold = True

            for i, miss in enumerate(missing, start=1):
                row = fill_tbl.rows[i].cells
                row[0].text = miss
                row[1].text = it.get("suggestion", "") if i == 1 else ""
                row[2].text = ""

    # ── 輸出 bytes ──
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ─── PDF 報告（簡化版 HTML→PDF）────────────────────────────────────

def generate_html_report(check_result: dict) -> str:
    """產出 HTML 格式的報告，可直接在瀏覽器列印成 PDF。"""
    score   = check_result.get("score", 0)
    passed  = check_result.get("passed", False)
    scene   = check_result.get("scene", "")
    level   = check_result.get("level", "")
    items   = check_result.get("items", [])
    stats   = check_result.get("stats", {})
    comment = check_result.get("overall_comment", "")

    rows_html = ""
    for it in items:
        result = it.get("result", "")
        color_map = {
            "已完成":"#d4edda","部分完成":"#fff3cd","未完成":"#f8d7da",
            "需補件":"#d1ecf1","需確認":"#e2d9f3","不適用":"#e9ecef"
        }
        bg = color_map.get(result, "#fff")
        missing_html = ""
        if it.get("missing"):
            missing_html = "<br><b>需補充：</b>" + "、".join(it["missing"])
        sug_html = f"<br><b>建議：</b>{it['suggestion']}" if it.get("suggestion") else ""
        rows_html += f"""
        <tr style="background:{bg}">
          <td>{it.get('seq','')}</td>
          <td>{it.get('item','')}</td>
          <td><b>{result}</b></td>
          <td>{it.get('evidence','')}{missing_html}{sug_html}</td>
        </tr>"""

    stat_html = " | ".join(f"{k}：{v}" for k, v in stats.items())

    return f"""<!DOCTYPE html>
<html lang="zh-Hant"><head>
<meta charset="utf-8">
<title>專案檢核報告</title>
<style>
  body {{ font-family: "Noto Sans TC", sans-serif; margin:32px; color:#222; }}
  h1 {{ color:#1a237e; }} h2 {{ color:#283593; border-bottom:2px solid #e8eaf6; }}
  .badge {{ display:inline-block; padding:4px 12px; border-radius:20px; font-weight:bold; }}
  .pass {{ background:#d4edda; color:#155724; }} .fail {{ background:#f8d7da; color:#721c24; }}
  table {{ width:100%; border-collapse:collapse; margin:16px 0; }}
  th {{ background:#1a237e; color:#fff; padding:8px; text-align:left; }}
  td {{ padding:8px; border:1px solid #dee2e6; vertical-align:top; }}
  .meta {{ background:#f8f9fa; padding:12px; border-radius:8px; margin:16px 0; }}
  @media print {{ .no-print {{ display:none; }} }}
</style></head><body>
<h1>📋 專案規劃檢核報告</h1>
<div class="meta">
  <b>場景：</b>{scene} &nbsp;|&nbsp;
  <b>層級：</b>{level} 級 &nbsp;|&nbsp;
  <b>Member：</b>{check_result.get('member_name','—')} &nbsp;|&nbsp;
  <b>日期：</b>{check_result.get('timestamp','')[:10]}
</div>

<h2>TC4 評分</h2>
<p style="font-size:2em;margin:0"><b>{score}</b> / 100 &nbsp;
  <span class="badge {'pass' if passed else 'fail'}">
    {'✅ 可找 William 開會' if passed else '❌ 未達開會門檻'}</span></p>
<p>{stat_html}</p>

{f'<h2>整體評語</h2><p>{comment}</p>' if comment else ''}

<h2>逐項檢核結果</h2>
<table>
  <thead><tr><th>#</th><th>檢核項目</th><th>結果</th><th>說明</th></tr></thead>
  <tbody>{rows_html}</tbody>
</table>

<div class="no-print" style="margin-top:32px">
  <button onclick="window.print()">🖨️ 列印 / 儲存 PDF</button>
</div>
</body></html>"""
