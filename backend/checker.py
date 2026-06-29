# -*- coding: utf-8 -*-
"""
checker.py — 核心檢核邏輯
"""

import io
import uuid
from pathlib import Path
from datetime import datetime

from docx import Document as DocxDocument
import openpyxl

import gemini_client
from db_loader import filter_items, detect_scene, extract_skill_context, get_mention_weight
from config import TC4_THRESHOLD, LEVEL_ORDER, IMPORTANCE_ORDER


# --- 文件解析 ---

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = "  |  ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def extract_text_from_xlsx(file_bytes: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"[工作表：{ws.title}]")
        for row in ws.iter_rows(values_only=True):
            row_text = "  ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                lines.append(row_text)
    wb.close()
    return "\n".join(lines)


def extract_text(file_bytes: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        return extract_text_from_docx(file_bytes)
    elif suffix == ".xlsx":
        return extract_text_from_xlsx(file_bytes)
    else:
        raise ValueError(f"不支援的檔案格式：{suffix}")


# --- TC4 評分（扣分制）---
#
# 起點 100 分。每個被檢核項目（不適用者排除）依重要性有基礎扣分點：
#   紅 = 3 點、黃 = 2 點、白 = 1 點
# 再乘以 mention_weight（高頻提到的項目權重更高）。
#
# 結果對應扣分比例：
#   已完成  → 0%（不扣）
#   部分完成 → 50% 扣
#   需補件  → 33% 扣
#   需確認  → 25% 扣
#   未完成  → 100% 扣
#   不適用  → 排除（不計入分子與分母）
#
# 最終分數 = 100 - (實際扣分總點數 / 最大可能扣分總點數) × 100
# 這等同於：score = (已得分點數 / 最高可得點數) × 100

DEDUCT_RATIO = {
    "已完成":   0.0,
    "部分完成": 0.5,
    "需補件":   1/3,
    "需確認":   0.25,
    "未完成":   1.0,
}

# 重要性對應基礎扣分點（影響未完成扣多少）
IMPORTANCE_WEIGHT = {
    "紅": 3.0,
    "黃": 2.0,
    "白": 1.0,
}


def calculate_tc4_score(items_result: list, db_items: list) -> dict:
    db_map = {it["seq"]: it for it in db_items}
    total_max_points = 0.0  # 最大可能扣分點
    total_deducted   = 0.0  # 實際扣分點
    detail = []

    for res in items_result:
        seq    = res.get("seq")
        result = res.get("result", "未完成")

        # 不適用 → 完全排除
        if result == "不適用":
            continue

        db_item = db_map.get(seq, {})
        importance = db_item.get("importance", "白")
        mention    = db_item.get("mention_count", 0)

        imp_w  = IMPORTANCE_WEIGHT.get(importance, 1.0)
        men_w  = get_mention_weight(mention)
        points = imp_w * men_w           # 此項目最大扣分點

        ratio    = DEDUCT_RATIO.get(result, 1.0)
        deducted = points * ratio

        total_max_points += points
        total_deducted   += deducted

        detail.append({
            "seq":        seq,
            "item":       res.get("item", ""),
            "result":     result,
            "importance": importance,
            "max_points": round(points, 2),
            "deducted":   round(deducted, 2),
        })

    if total_max_points == 0:
        final_score = 0.0
    else:
        final_score = round(100 - (total_deducted / total_max_points) * 100, 1)

    return {
        "score":  final_score,
        "passed": final_score >= TC4_THRESHOLD,
        "detail": detail,
    }


# --- 主流程 ---

def run_check(
    file_bytes: bytes,
    filename: str,
    scene_hint: str = "",
    member_name: str = "",
    context_input: str = "",
    gemini_api_key: str = None,
    gemini_model: str = None,
    pre_extracted: str = None,
    api_provider: str = "gemini",
    third_party_key: str = None,
    third_party_model: str = None,
    third_party_base_url: str = "https://api.runapi.sbs/v1",
) -> dict:
    check_id  = str(uuid.uuid4())
    timestamp = datetime.now().isoformat()

    # 1. 解析文件
    doc_text = ""
    if pre_extracted:
        doc_text = pre_extracted
    elif file_bytes and filename:
        doc_text = extract_text(file_bytes, filename)

    if context_input.strip():
        context_section = f"【學員說明與背景】\n{context_input.strip()}\n\n"
        doc_text = context_section + doc_text

    detect_text = doc_text if doc_text else context_input

    # 2. 場景辨識（第一輪 AI）
    if scene_hint:
        # 使用者手動指定場景 → 直接用，仍跑關鍵字判斷層級
        keyword_result = detect_scene(detect_text)
        scenes = [scene_hint, "G. 專案管理基本功"]
        if "G. 專案管理基本功" not in scenes:
            scenes.append("G. 專案管理基本功")
        level      = keyword_result["level"]
        confidence = 100
        scene_note = f"手動指定：{scene_hint}"
    else:
        # AI 第一輪：完整閱讀文件，判斷主/次場景 + 層級
        ai_scene_result = gemini_client.detect_scenes_ai(
            doc_text              = detect_text,
            api_key               = gemini_api_key,
            provider              = api_provider,
            third_party_key       = third_party_key,
            third_party_model     = third_party_model,
            third_party_base_url  = third_party_base_url,
        )
        scenes     = ai_scene_result.get("_scenes", [])
        level      = ai_scene_result.get("level", "B")
        confidence = 85   # AI 辨識，給固定信心值
        scene_note = ai_scene_result.get("level_reason", "")

        # AI 失敗 fallback → 關鍵字比對
        if not scenes:
            kw = detect_scene(detect_text)
            scenes     = [kw["scene"], "G. 專案管理基本功"]
            level      = kw["level"]
            confidence = kw["confidence"]
            scene_note = "AI 辨識失敗，改用關鍵字比對"

    # 主場景（第一個）供顯示用
    primary_scene = scenes[0] if scenes else "G. 專案管理基本功"

    # 3. 篩選檢核項目 + skill context
    db_items      = filter_items(scenes, level)
    skill_context = extract_skill_context(scenes, level)

    # 4. 呼叫 AI 第二輪（逐項檢核）
    ai_result = gemini_client.analyze(
        doc_text             = doc_text,
        scene                = scenes,
        level                = level,
        check_items          = db_items,
        skill_context        = skill_context,
        api_key              = gemini_api_key,
        model_override       = gemini_model,
        provider             = api_provider,
        third_party_key      = third_party_key,
        third_party_model    = third_party_model,
        third_party_base_url = third_party_base_url,
    )

    if "error" in ai_result:
        return {
            "check_id":  check_id,
            "error":     ai_result["error"],
            "timestamp": timestamp,
        }

    # 5. TC4 評分（扣分制）
    tc4 = calculate_tc4_score(ai_result.get("items", []), db_items)

    # 6. 統計
    items = ai_result.get("items", [])
    stats = {r: sum(1 for i in items if i.get("result") == r)
             for r in ["已完成", "部分完成", "未完成", "不適用", "需補件", "需確認"]}

    return {
        "check_id":          check_id,
        "timestamp":         timestamp,
        "filename":          filename,
        "member_name":       member_name,
        "scene":             ai_result.get("scene_confirmed", primary_scene),
        "scenes":            scenes,
        "level":             ai_result.get("level_confirmed", level),
        "scene_note":        scene_note,
        "auto_confidence":   confidence,
        "score":             tc4["score"],
        "passed":            tc4["passed"],
        "stats":             stats,
        "items":             items,
        "overall_comment":   ai_result.get("overall_comment", ""),
        "total_db_items":    len(db_items),
        "has_context_input": bool(context_input.strip()),
    }
